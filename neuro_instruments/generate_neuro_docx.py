#!/usr/bin/env python3
"""
Generator: creates one DOCX per category for neuro instruments,
matching the format of the veterinary instrument example.
"""

import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image as PILImage

BLUE   = RGBColor(0x1F, 0x4E, 0x9C)
ORANGE = RGBColor(0xE8, 0x73, 0x0C)
GRAY   = RGBColor(0x8A, 0x7F, 0x6A)
BLACK  = RGBColor(0x00, 0x00, 0x00)

SZ_HEADER   = 40
SZ_CATEGORY = 28
SZ_INSTNAME = 24
SZ_HEADING  = 22
SZ_BODY     = 21

IMAGE_WIDTH_CM = 5.5
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
IMG_BASE = os.path.join(BASE_DIR, "INSTRUMENTS")
OUT_DIR  = os.path.join(BASE_DIR, "output_docx")
os.makedirs(OUT_DIR, exist_ok=True)


def _set_run_format(run, size_hp, bold=False, color=None):
    rpr = run._r.get_or_add_rPr()
    sz = OxmlElement('w:sz');   sz.set(qn('w:val'),   str(size_hp)); rpr.append(sz)
    szcs = OxmlElement('w:szCs'); szcs.set(qn('w:val'), str(size_hp)); rpr.append(szcs)
    if bold:
        rpr.append(OxmlElement('w:b')); rpr.append(OxmlElement('w:bCs'))
    if color:
        c = OxmlElement('w:color')
        c.set(qn('w:val'), '{:02X}{:02X}{:02X}'.format(color[0], color[1], color[2]))
        rpr.append(c)


def _set_para_spacing(para, before=0, after=0):
    ppr = para._p.get_or_add_pPr()
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:before'), str(before)); sp.set(qn('w:after'), str(after))
    ppr.append(sp)


def _set_para_align(para, align='center'):
    ppr = para._p.get_or_add_pPr()
    jc = OxmlElement('w:jc'); jc.set(qn('w:val'), align); ppr.append(jc)


def add_styled_para(doc, text, size_hp, bold=False, color=None, align=None, before=0, after=0):
    para = doc.add_paragraph()
    _set_para_spacing(para, before=before, after=after)
    if align:
        _set_para_align(para, align)
    run = para.add_run(text)
    _set_run_format(run, size_hp, bold=bold, color=color)
    return para


def set_cell_borders_none(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:val'), 'none'); el.set(qn('w:color'), 'FFFFFF'); el.set(qn('w:sz'), '0')
        borders.append(el)
    tcpr.append(borders)


def set_cell_margins(cell, top=0, left=0, bottom=0, right=0):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement('w:tcMar')
    for side, val in [('top', top), ('left', left), ('bottom', bottom), ('right', right)]:
        el = OxmlElement(f'w:{side}'); el.set(qn('w:type'), 'dxa'); el.set(qn('w:w'), str(val))
        mar.append(el)
    tcpr.append(mar)


def set_cell_width(cell, width_twips):
    tcpr = cell._tc.get_or_add_tcPr()
    w = OxmlElement('w:tcW'); w.set(qn('w:type'), 'dxa'); w.set(qn('w:w'), str(width_twips))
    tcpr.append(w)


def set_table_width(table, width_twips):
    tbl = table._tbl
    tblpr = tbl.find(qn('w:tblPr'))
    if tblpr is None:
        tblpr = OxmlElement('w:tblPr'); tbl.insert(0, tblpr)
    w = OxmlElement('w:tblW'); w.set(qn('w:type'), 'dxa'); w.set(qn('w:w'), str(width_twips))
    tblpr.append(w)


def add_table_borders(table):
    tbl = table._tbl
    tblpr = tbl.find(qn('w:tblPr'))
    if tblpr is None:
        tblpr = OxmlElement('w:tblPr'); tbl.insert(0, tblpr)
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:val'), 'single'); el.set(qn('w:color'), 'auto'); el.set(qn('w:sz'), '4')
        borders.append(el)
    tblpr.append(borders)


def insert_image_in_cell(cell, img_path):
    para = cell.paragraphs[0]
    run = para.add_run()
    try:
        with PILImage.open(img_path) as im:
            w_px, h_px = im.size
        run.add_picture(img_path, width=Cm(IMAGE_WIDTH_CM), height=Cm(IMAGE_WIDTH_CM * h_px / w_px))
    except Exception:
        para.add_run(f'[Image: {os.path.basename(img_path)}]')


def add_instrument_block(doc, category_display, instrument, cat_folder):
    name   = instrument['name']
    img_fn = instrument['image']
    i_use  = instrument['intended_use']
    func   = instrument['function']
    advs   = instrument['advantages']
    faqs   = instrument['faqs']
    img_path = os.path.join(IMG_BASE, cat_folder, img_fn)

    add_styled_para(doc, 'NEURO INSTRUMENTS', SZ_HEADER, bold=True, color=BLUE, align='center', after=60)
    add_styled_para(doc, category_display, SZ_CATEGORY, bold=True, color=ORANGE, align='center', after=160)
    add_styled_para(doc, name, SZ_INSTNAME, bold=True, color=GRAY, after=120)
    add_styled_para(doc, 'Intended Use', SZ_HEADING, bold=True, color=BLACK, before=160, after=80)

    table = doc.add_table(rows=1, cols=2)
    set_table_width(table, 9360)
    add_table_borders(table)
    left_cell  = table.cell(0, 0)
    right_cell = table.cell(0, 1)
    set_cell_width(left_cell,  5560)
    set_cell_width(right_cell, 3800)
    set_cell_borders_none(left_cell)
    set_cell_borders_none(right_cell)
    set_cell_margins(left_cell,  right=160)
    set_cell_margins(right_cell, left=80)

    lp = left_cell.paragraphs[0]
    _set_para_align(lp, 'both')
    lr = lp.add_run(i_use)
    _set_run_format(lr, SZ_BODY)

    if os.path.exists(img_path):
        insert_image_in_cell(right_cell, img_path)
    else:
        right_cell.paragraphs[0].add_run(f'[Missing: {img_fn}]')

    add_styled_para(doc, 'Function', SZ_HEADING, bold=True, color=BLACK, before=160, after=80)
    p = doc.add_paragraph()
    _set_run_format(p.add_run(func), SZ_BODY)

    add_styled_para(doc, 'Advantages', SZ_HEADING, bold=True, color=BLACK, before=160, after=80)
    for title, body in advs:
        p = doc.add_paragraph()
        _set_run_format(p.add_run(f'{title}: '), SZ_BODY, bold=True)
        _set_run_format(p.add_run(body), SZ_BODY, bold=True)

    add_styled_para(doc, 'FAQ', SZ_HEADING, bold=True, color=BLACK, before=160, after=80)
    for q, a in faqs:
        pq = doc.add_paragraph()
        _set_run_format(pq.add_run(f'Q: {q}'), SZ_BODY, bold=True)
        pa = doc.add_paragraph()
        _set_run_format(pa.add_run(f'A: {a}'), SZ_BODY)


def generate_category_docx(cat_key, cat_data):
    display     = cat_data.get('display_name', cat_key)
    instruments = cat_data['instruments']
    doc = Document()
    for i, instr in enumerate(instruments):
        add_instrument_block(doc, display, instr, cat_key)
        if i < len(instruments) - 1:
            doc.add_page_break()
    safe_name = display.replace('/', '-').replace('&', 'and').replace(',', '')
    out_path = os.path.join(OUT_DIR, f'{safe_name}.docx')
    doc.save(out_path)
    print(f'  Saved: {out_path}')
    return out_path


def main():
    from content_data_neuro import NEURO_CATEGORIES
    print(f'Generating {len(NEURO_CATEGORIES)} DOCX file(s)...\n')
    for cat_key, cat_data in NEURO_CATEGORIES.items():
        print(f'Processing: {cat_key}')
        try:
            generate_category_docx(cat_key, cat_data)
        except Exception as e:
            print(f'  ERROR: {e}')
    print('\nDone.')


if __name__ == '__main__':
    main()
